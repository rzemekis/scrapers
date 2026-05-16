import time
import os
import tempfile
import shutil
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC


def run_capture_grabber():
    file_path = 'LEGO.xlsx'

    try:
        wb = load_workbook(file_path)
        ws = wb.active
    except FileNotFoundError:
        print(f"Файл {file_path} не найден!")
        return

    temp_dir = tempfile.mkdtemp()

    options = webdriver.ChromeOptions()
    options.add_argument(
        "user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36")
    options.add_experimental_option("excludeSwitches", ["enable-automation"])

    driver = webdriver.Chrome(options=options)
    driver.maximize_window()

    try:
        driver.get("https://www.bricklink.com/v2/main.page")
        print("\n>>> 1. Решите капчу и примите куки.")
        print(">>> 2. Нажмите ENTER в консоли для начала.\n")
        input()

        count_saved = 0

        for row in ws.iter_rows(min_row=1, max_col=2):
            cell_a = row[0]
            cell_b = row[1]

            if not cell_a.value or cell_b.value:
                continue

            article = str(cell_a.value).split('.')[0].strip()
            print(f"[{article}] >>", end=" ")

            try:
                target_url = f"https://www.bricklink.com/v2/catalog/catalogitem.page?S={article}-1"
                driver.get(target_url)

                main_img_el = WebDriverWait(driver, 7).until(
                    EC.visibility_of_element_located((By.ID, "_idImageMain"))
                )

                time.sleep(1)

                img_binary = main_img_el.screenshot_as_png
                img_data = BytesIO(img_binary)

                with PILImage.open(img_data) as pil_img:
                    if pil_img.mode in ("RGBA", "P"):
                        pil_img = pil_img.convert("RGB")

                    pil_img.thumbnail((150, 150))

                    temp_img_path = os.path.join(temp_dir, f"{article}.png")
                    pil_img.save(temp_img_path)

                    xl_img = XLImage(temp_img_path)
                    ws.add_image(xl_img, f'B{cell_a.row}')
                    ws.row_dimensions[cell_a.row].height = 115

                    print("СКОПИРОВАНО.")
                    count_saved += 1

                    if count_saved % 10 == 0:
                        wb.save(file_path)

            except Exception:
                print("НЕ НАЙДЕНО.")
                continue

        ws.column_dimensions['B'].width = 25
        wb.save(file_path)
        print(f"\nГОТОВО! Добавлено: {count_saved}")

    finally:
        driver.quit()
        if os.path.exists(temp_dir):
            shutil.rmtree(temp_dir)


if __name__ == "__main__":
    run_capture_grabber()